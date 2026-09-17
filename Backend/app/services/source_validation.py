import json
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

_ROOT = Path(__file__).resolve().parent.parent.parent / "data" / "validations" / "sources"


def load_rules(source: str) -> dict:
    path = _ROOT / source / "rules.json"
    if not path.exists():
        raise HTTPException(status_code=500, detail=f"No validation rules for source {source}")
    return json.loads(path.read_text(encoding="utf-8"))


def rules_dir(source: str) -> Path:
    return _ROOT / source


@dataclass
class Context:
    filename: str
    payload: bytes
    rules: dict
    book: object | None = None
    sheet: Worksheet | None = None
    headers: dict[str, int] = field(default_factory=dict)
    rows: list[dict] = field(default_factory=list)
    header_cells: list[str] = field(default_factory=list)
    issue_number: int | None = None


def _norm(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value).strip()


def _cell(value: object) -> str:
    text = _norm(value)
    return text.upper() if text.upper() in {"Y", "OUT"} else text


def _format_value(column: dict, raw: object) -> str:
    text = _cell(raw)
    compact = text.upper().replace(" ", "")
    if compact in {"N/A", "NA"}:
        return "N/A"
    if column.get("id") == "type":
        kinds = {
            "SERVICE BULLETIN": "Service bulletin",
            "SB": "Service bulletin",
            "MAINTENANCE TASK": "Maintenance task",
            "MT": "Maintenance task",
        }
        return kinds.get(text.upper(), text)
    pad = column.get("pad")
    if pad and text:
        try:
            return f"{int(float(text)):0{int(pad)}d}"
        except ValueError:
            return text
    return text


def _header_names(column: dict) -> list[str]:
    names = [column["header"], *column.get("aliases", [])]
    return [_norm(name).upper() for name in names]


def _open_book(payload: bytes):
    return load_workbook(BytesIO(payload), data_only=True)


_VERSION_HEADER = re.compile(r"^(\d{1,2})\.(\d{2})$")


def _version_id(header: str) -> str | None:
    match = _VERSION_HEADER.match(_norm(header))
    if not match:
        return None
    return f"{int(match.group(1)):02d}.{match.group(2)}"


def _version_spec(rules: dict, col_id: str) -> dict:
    spec = dict(rules.get("version") or {})
    spec["id"] = col_id
    spec["header"] = col_id
    return spec


def _column_def(rules: dict, col_id: str) -> dict:
    for column in rules["columns"]:
        if column["id"] == col_id:
            return column
    return _version_spec(rules, col_id)


def _named_ids(rules: dict) -> set[str]:
    return {column["id"] for column in rules["columns"]}


def _map_headers(sheet: Worksheet, rules: dict) -> dict[str, int]:
    header_row = int(rules.get("sheet", {}).get("header_row") or 1)
    cells = [_norm(cell).upper() for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    mapped: dict[str, int] = {}
    missing: list[str] = []
    for column in rules["columns"]:
        found = None
        for name in _header_names(column):
            if name in cells:
                found = cells.index(name)
                break
        if found is None:
            if column.get("required", True):
                missing.append(column["header"])
            continue
        mapped[column["id"]] = found
    if missing:
        raise ValueError("Missing columns: " + ", ".join(missing))
    used = set(mapped.values())
    last_index = next((index for index in range(len(cells) - 1, -1, -1) if cells[index]), None)
    for index, name in enumerate(cells):
        if index == last_index or index in used or not name:
            continue
        vid = _version_id(name)
        if not vid or vid in mapped:
            continue
        mapped[vid] = index
        used.add(index)
    return mapped


def _read_rows(sheet: Worksheet, rules: dict, headers: dict[str, int]) -> list[dict]:
    header_row = int(rules.get("sheet", {}).get("header_row") or 1)
    rows: list[dict] = []
    for index, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not values or all(value is None or str(value).strip() == "" for value in values):
            continue
        row = {"_line": index}
        for col_id, col_index in headers.items():
            raw = values[col_index] if col_index < len(values) else None
            row[col_id] = _format_value(_column_def(rules, col_id), raw)
        rows.append(row)
    if not rows:
        raise ValueError("The sheet has no data rows")
    return rows


def _row_columns(rules: dict, rows: list[dict]) -> list[dict]:
    named = list(rules["columns"])
    known = _named_ids(rules)
    extra = []
    if rows:
        for key in rows[0]:
            if key in known or key == "_line":
                continue
            extra.append(_version_spec(rules, key))
    return named + extra


def _empty_errors(rules: dict, rows: list[dict]) -> list[str]:
    errors: list[str] = []
    for column in _row_columns(rules, rows):
        if column.get("allow_empty", False):
            continue
        for row in rows:
            if not row.get(column["id"]):
                errors.append(f"Row {row['_line']}: {column['header']} is empty")
                if len(errors) >= 3:
                    return errors
    return errors


def _value_errors(rules: dict, rows: list[dict]) -> list[str]:
    errors: list[str] = []
    for column in _row_columns(rules, rows):
        pattern = column.get("pattern")
        allowed = column.get("enum")
        min_length = column.get("min_length")
        compiled = re.compile(pattern) if pattern else None
        for row in rows:
            value = row.get(column["id"]) or ""
            if not value:
                continue
            if compiled and not compiled.match(value):
                errors.append(f"Row {row['_line']}: {column['header']} has an invalid format")
            elif allowed and value not in allowed:
                errors.append(f"Row {row['_line']}: {column['header']} must be {' / '.join(allowed)}")
            elif min_length and len(value) < min_length:
                errors.append(f"Row {row['_line']}: {column['header']} is too short")
            if len(errors) >= 3:
                return errors
    return errors


def _version_ids(rules: dict, rows: list[dict]) -> list[str]:
    known = _named_ids(rules)
    if not rows:
        return []
    return [key for key in rows[0] if key not in known and key != "_line"]


def _flag_errors(rules: dict, rows: list[dict]) -> list[str]:
    versions = _version_ids(rules, rows)
    errors: list[str] = []
    for row in rows:
        locked = False
        for col in versions:
            value = row.get(col) or ""
            if locked and value == "Y":
                errors.append(f"Row {row['_line']}: {col} cannot be Y after OUT")
            if value == "OUT":
                locked = True
            if len(errors) >= 3:
                return errors
    return errors


def _unique_errors(rules: dict, rows: list[dict]) -> list[str]:
    errors: list[str] = []
    for column in rules["columns"]:
        if not column.get("unique"):
            continue
        seen: dict[str, int] = {}
        for row in rows:
            keys = [column["id"], *column.get("unique_with", [])]
            value = " · ".join(row.get(key) or "" for key in keys)
            if not (row.get(column["id"]) or ""):
                continue
            if value in seen:
                errors.append(f"Row {row['_line']}: duplicated {column['header']}")
                if len(errors) >= 3:
                    return errors
            else:
                seen[value] = row["_line"]
    return errors


_ISSUE_IN_NAME = re.compile(r"ISSUE[\s_\-]*(\d{1,2}|XX)\s*$", re.I)


def _header_key(value: str) -> str:
    return _norm(value).upper().replace(" ", "")


def _sheet_headers(sheet: Worksheet, rules: dict) -> list[str]:
    header_row = int(rules.get("sheet", {}).get("header_row") or 1)
    cells = [_norm(cell) for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    while cells and not cells[-1]:
        cells.pop()
    if not cells:
        raise ValueError("The sheet has no header row")
    return cells


def _issue_from_last_column(headers: list[str]) -> int:
    last = headers[-1] if headers else ""
    match = _ISSUE_IN_NAME.search(_norm(last))
    if not match:
        found = last or "(empty)"
        raise ValueError(f"Last column must contain Issue XX, found '{found}'")
    token = match.group(1)
    if token.upper() == "XX":
        nums = [
            int(vid[:2])
            for name in headers[:-1]
            if (vid := _version_id(name)) and vid.endswith(".00")
        ]
        if not nums:
            raise ValueError("Last column is Issue XX but no NN.00 column was found")
        return max(nums)
    return int(token)


def run_check(check_id: str, ctx: Context) -> tuple[bool, str]:
    rules = ctx.rules
    file_rules = rules.get("file") or {}
    if check_id == "source":
        if not ctx.payload:
            return False, "The uploaded file is empty"
        max_bytes = int(file_rules.get("max_bytes") or 0)
        if max_bytes and len(ctx.payload) > max_bytes:
            return False, "The file exceeds the maximum size"
        return True, "File received"
    if check_id == "format":
        ext = Path(ctx.filename).suffix.lower()
        allowed = {item.lower() for item in file_rules.get("extensions") or []}
        if ext not in allowed:
            return False, f"{ext or 'Unknown type'} is not a valid {rules['label']} Excel"
        return True, "Excel format accepted"
    if check_id == "workbook":
        try:
            ctx.book = _open_book(ctx.payload)
            wanted = rules.get("sheet", {}).get("name")
            if wanted:
                if wanted not in ctx.book.sheetnames:
                    return False, f"Sheet '{wanted}' was not found"
                ctx.sheet = ctx.book[wanted]
            else:
                ctx.sheet = ctx.book.active
        except Exception as exc:  # noqa: BLE001
            return False, str(exc) or "The workbook could not be opened"
        return True, "Workbook opened"
    if check_id == "issue_name":
        try:
            ctx.header_cells = _sheet_headers(ctx.sheet, rules)
            ctx.issue_number = _issue_from_last_column(ctx.header_cells)
        except ValueError as exc:
            return False, str(exc)
        return True, f"Last column is Issue {ctx.issue_number:02d}"
    if check_id == "issue_version":
        if ctx.issue_number is None:
            return False, "Issue number was not found in the last column"
        wanted = f"{ctx.issue_number:02d}.00"
        keys = {_header_key(name) for name in ctx.header_cells}
        aliases = {_header_key(wanted), _header_key(f"{ctx.issue_number}.00")}
        if not keys & aliases:
            return False, f"Column {wanted} is required when the last column is Issue {ctx.issue_number:02d}"
        return True, f"Column {wanted} matches Issue {ctx.issue_number:02d}"
    if check_id == "headers":
        try:
            ctx.headers = _map_headers(ctx.sheet, rules)
        except Exception as exc:  # noqa: BLE001
            return False, str(exc)
        return True, "Required columns found"
    if check_id == "empty":
        try:
            ctx.rows = _read_rows(ctx.sheet, rules, ctx.headers)
        except Exception as exc:  # noqa: BLE001
            return False, str(exc)
        errors = _empty_errors(rules, ctx.rows)
        return (False, errors[0]) if errors else (True, "No empty required cells")
    if check_id == "values":
        errors = _value_errors(rules, ctx.rows)
        return (False, errors[0]) if errors else (True, "Column formats are valid")
    if check_id == "flags":
        errors = _flag_errors(rules, ctx.rows)
        return (False, errors[0]) if errors else (True, "Flags are sequential")
    if check_id == "unique":
        errors = _unique_errors(rules, ctx.rows)
        return (False, errors[0]) if errors else (True, "References are unique")
    if check_id == "store":
        return (True, "Ready to store") if ctx.rows else (False, "No rows to store")
    return False, f"Unknown check {check_id}"


def to_pd_rows(rows: list[dict], rules: dict) -> tuple[list[str], list[dict]]:
    versions = _version_ids(rules, rows)
    parsed = []
    for row in rows:
        parsed.append(
            {
                "item_ref": row["item_ref"],
                "reference": row["reference"],
                "revision": row["revision"],
                "ata": row["ata"],
                "description": row.get("description") or row.get("title") or "",
                "type": row.get("type") or "",
                "source_material": row.get("source_material") or "",
                "source_hours": row.get("source_hours") or "",
                "fin_position": row.get("fin_position") or "",
                "pn": row.get("pn") or "",
                "sn": row.get("sn") or "",
                "flags": {col: row[col] for col in versions},
            }
        )
    return versions, parsed


def to_named_rows(rows: list[dict], rules: dict) -> list[dict]:
    ids = [column["id"] for column in rules["columns"]]
    return [{col: row.get(col) or "" for col in ids} for row in rows]


def write_sample_workbook(path: Path, rows: list[dict]) -> None:
    rules = load_rules("pd")
    book = Workbook()
    sheet = book.active
    sheet.title = "PD"
    versions = list(rows[0]["flags"].keys()) if rows else []
    columns = list(rules["columns"])
    issue_n = max((int(col[:2]) for col in versions if col.endswith(".00")), default=1)
    issue_header = f"Issue {issue_n:02d}"
    sheet.append([column["header"] for column in columns] + versions + [issue_header])
    for cell in sheet[1]:
        cell.number_format = "@"
    for row in rows:
        line = [row.get(column["id"], "") for column in columns]
        line.extend(row["flags"][col] for col in versions)
        line.append("")
        sheet.append(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
