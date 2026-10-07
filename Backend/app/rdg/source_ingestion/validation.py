import json
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.paths import DATA_DIR

_ROOT = DATA_DIR / "validations" / "sources"


def load_rules(source: str) -> dict:
    path = _ROOT / source / "rules.json"
    if not path.exists():
        raise HTTPException(status_code=500, detail=f"No validation rules for source {source}")
    return json.loads(path.read_text(encoding="utf-8"))


def rules_dir(source: str) -> Path:
    return _ROOT / source


@dataclass
class Context:
    filename: str
    payload: bytes
    rules: dict
    book: object | None = None
    sheet: Worksheet | None = None
    headers: dict[str, int] = field(default_factory=dict)
    rows: list[dict] = field(default_factory=list)
    header_cells: list[str] = field(default_factory=list)
    issue_number: int | None = None
    file_version: str | None = None
    failed: bool = False
    fail_rows: list[dict] = field(default_factory=list)


def _norm(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value).strip()


def _cell(value: object) -> str:
    text = _norm(value)
    upper = text.upper()
    if upper in {"Y", "OUT", "UNDER REVIEW"}:
        return upper
    return text


def _format_value(column: dict, raw: object) -> str:
    text = _cell(raw)
    compact = text.upper().replace(" ", "")
    if compact in {"N/A", "NA"}:
        return "N/A"
    if column.get("id") == "type":
        kinds = {
            "SERVICE BULLETIN": "Service bulletin",
            "SB": "Service bulletin",
            "MAINTENANCE TASK": "Maintenance task",
            "MT": "Maintenance task",
        }
        return kinds.get(text.upper(), text)
    pad = column.get("pad")
    if pad and text:
        try:
            return f"{int(float(text)):0{int(pad)}d}"
        except ValueError:
            return text
    return text


def _header_names(column: dict) -> list[str]:
    names = [column["header"], *column.get("aliases", [])]
    return [_norm(name).upper() for name in names]


def _open_book(payload: bytes):
    return load_workbook(BytesIO(payload), data_only=True)


_VERSION_HEADER = re.compile(r"^(\d{1,2})\.(\d{2})$")


def _version_id(header: str) -> str | None:
    match = _VERSION_HEADER.match(_norm(header))
    if not match:
        return None
    return f"{int(match.group(1)):02d}.{match.group(2)}"


def _version_spec(rules: dict, col_id: str) -> dict:
    spec = dict(rules.get("version") or {})
    spec["id"] = col_id
    spec["header"] = col_id
    spec.setdefault("allow_empty", True)
    spec.pop("enum", None)
    spec.pop("pattern", None)
    return spec


def _column_def(rules: dict, col_id: str) -> dict:
    for column in rules["columns"]:
        if column["id"] == col_id:
            return column
    return _version_spec(rules, col_id)


def _named_ids(rules: dict) -> set[str]:
    return {column["id"] for column in rules["columns"]}


def _map_headers(sheet: Worksheet, rules: dict) -> dict[str, int]:
    header_row = int(rules.get("sheet", {}).get("header_row") or 1)
    cells = [_norm(cell).upper() for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    mapped: dict[str, int] = {}
    missing: list[str] = []
    for column in rules["columns"]:
        found = None
        for name in _header_names(column):
            if name in cells:
                found = cells.index(name)
                break
        if found is None:
            if column.get("required", True):
                missing.append(column["header"])
            continue
        mapped[column["id"]] = found
    if missing:
        raise ValueError("Missing columns: " + ", ".join(missing))
    used = set(mapped.values())
    for index, name in enumerate(cells):
        if index in used or not name:
            continue
        vid = _version_id(name)
        if not vid or vid in mapped:
            continue
        mapped[vid] = index
        used.add(index)
    return mapped


def _read_rows(sheet: Worksheet, rules: dict, headers: dict[str, int]) -> list[dict]:
    header_row = int(rules.get("sheet", {}).get("header_row") or 1)
    rows: list[dict] = []
    for index, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not values or all(value is None or str(value).strip() == "" for value in values):
            continue
        row = {"_line": index}
        for col_id, col_index in headers.items():
            raw = values[col_index] if col_index < len(values) else None
            row[col_id] = _format_value(_column_def(rules, col_id), raw)
        rows.append(row)
    if not rows:
        raise ValueError("The sheet has no data rows")
    return rows


def _row_columns(rules: dict, rows: list[dict]) -> list[dict]:
    named = list(rules["columns"])
    known = _named_ids(rules)
    extra = []
    if rows:
        for key in rows[0]:
            if key in known or key == "_line":
                continue
            extra.append(_version_spec(rules, key))
    return named + extra


def _empty_errors(rules: dict, rows: list[dict]) -> list[str]:
    errors: list[str] = []
    for column in _row_columns(rules, rows):
        if column.get("allow_empty", False):
            continue
        for row in rows:
            if not row.get(column["id"]):
                errors.append(f"Row {row['_line']}: {column['header']} is empty")
                if len(errors) >= 20:
                    return errors
    return errors


def _value_errors(rules: dict, rows: list[dict]) -> list[str]:
    errors: list[str] = []
    for column in _row_columns(rules, rows):
        pattern = column.get("pattern")
        allowed = column.get("enum")
        min_length = column.get("min_length")
        compiled = re.compile(pattern) if pattern else None
        for row in rows:
            value = row.get(column["id"]) or ""
            if not value:
                continue
            if compiled and not compiled.match(value):
                errors.append(f"Row {row['_line']}: {column['header']} has an invalid format")
            elif allowed and value not in allowed:
                errors.append(f"Row {row['_line']}: {column['header']} must be {' / '.join(allowed)}")
            elif min_length and len(value) < min_length:
                errors.append(f"Row {row['_line']}: {column['header']} is too short")
            if len(errors) >= 20:
                return errors
    return errors


def _version_ids(rules: dict, rows: list[dict]) -> list[str]:
    known = _named_ids(rules)
    if not rows:
        return []
    return [key for key in rows[0] if key not in known and key != "_line"]


def _flag_errors(rules: dict, rows: list[dict]) -> list[str]:
    versions = _version_ids(rules, rows)
    errors: list[str] = []
    for row in rows:
        locked = False
        for col in versions:
            value = row.get(col) or ""
            if locked and value == "Y":
                errors.append(f"Row {row['_line']}: {col} cannot be Y after OUT")
            if value == "OUT":
                locked = True
            if len(errors) >= 20:
                return errors
    return errors


def _unique_errors(rules: dict, rows: list[dict]) -> list[str]:
    errors: list[str] = []
    for column in rules["columns"]:
        if not column.get("unique"):
            continue
        seen: dict[str, int] = {}
        for row in rows:
            keys = [column["id"], *column.get("unique_with", [])]
            value = " · ".join(row.get(key) or "" for key in keys)
            if not (row.get(column["id"]) or ""):
                continue
            if value in seen:
                errors.append(f"Row {row['_line']}: duplicated {column['header']}")
                if len(errors) >= 20:
                    return errors
            else:
                seen[value] = row["_line"]
    return errors


_LINE_KEYS = ("task_reference", "revision", "fin_position", "pn", "sn")


def _joined(errors: list[str]) -> str:
    return "\n".join(errors)


def _fail_snap(row: dict) -> dict:
    snap = {"line": row["_line"]}
    for key, value in row.items():
        if key != "_line":
            snap[key] = value
    return snap


def _attach_fail_rows(ctx: Context, detail: str) -> None:
    if not ctx.rows or not detail:
        return
    lines: set[int] = set()
    for part in detail.split("\n"):
        head = part.split(":", 1)[0]
        if "row" not in head.lower():
            continue
        for match in re.finditer(r"\d+", head):
            lines.add(int(match.group(0)))
    seen = {item.get("line") for item in ctx.fail_rows}
    for row in ctx.rows:
        if row["_line"] in lines and row["_line"] not in seen:
            ctx.fail_rows.append(_fail_snap(row))
            seen.add(row["_line"])


def _fail_detail(ctx: Context, errors: list[str], ok_msg: str) -> tuple[bool, str]:
    if not errors:
        return True, ok_msg
    detail = _joined(errors)
    _attach_fail_rows(ctx, detail)
    return False, detail


def _line_unique_errors(rows: list[dict], version: str | None) -> tuple[list[str], list[dict]]:
    groups: dict[str, list[dict]] = {}
    issue = version or "issue"
    for row in rows:
        if not (row.get("task_reference") or "").strip():
            continue
        flag = (row.get(issue) or "").strip().upper() if version else ""
        value = " · ".join((row.get(key) or "").strip().upper() for key in _LINE_KEYS)
        groups.setdefault(f"{issue}|{flag}|{value}", []).append(row)
    errors: list[str] = []
    fail_rows: list[dict] = []
    for bunch in groups.values():
        if len(bunch) < 2:
            continue
        sample = bunch[0]
        listed = ", ".join(str(row["_line"]) for row in bunch)
        flag = (sample.get(issue) or "—") if version else "—"
        errors.append(
            f"Rows {listed}: duplicated TASK REFERENCE={sample.get('task_reference') or '—'} · "
            f"REVISION={sample.get('revision') or '—'} · FIN={sample.get('fin_position') or '—'} · "
            f"PNR={sample.get('pn') or '—'} · SNR={sample.get('sn') or '—'} · {issue}={flag}"
        )
        for row in bunch:
            fail_rows.append(_fail_snap(row))
    return errors, fail_rows


_FILE_VERSION = re.compile(r"(?<!\d)(\d{2}\.\d{2})(?!\d)")


def _header_key(value: str) -> str:
    return _norm(value).upper().replace(" ", "")


def _sheet_headers(sheet: Worksheet, rules: dict) -> list[str]:
    header_row = int(rules.get("sheet", {}).get("header_row") or 1)
    cells = [_norm(cell) for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    while cells and not cells[-1]:
        cells.pop()
    if not cells:
        raise ValueError("The sheet has no header row")
    return cells


def _version_from_filename(filename: str) -> str:
    matches = _FILE_VERSION.findall(Path(filename).stem)
    if not matches:
        raise ValueError("Filename must contain XX.XX (two digits, a dot, two digits)")
    return matches[-1]


def fill_issue_label(label: str, version: str | None) -> str:
    return label.replace("XX.XX", version) if version else label


def run_check(check_id: str, ctx: Context) -> tuple[bool, str]:
    rules = ctx.rules
    file_rules = rules.get("file") or {}
    if check_id == "source":
        if not ctx.payload:
            return False, "The uploaded file is empty"
        max_bytes = int(file_rules.get("max_bytes") or 0)
        if max_bytes and len(ctx.payload) > max_bytes:
            return False, "The file exceeds the maximum size"
        return True, "File received"
    if check_id == "format":
        ext = Path(ctx.filename).suffix.lower()
        allowed = {item.lower() for item in file_rules.get("extensions") or []}
        if ext not in allowed:
            return False, f"{ext or 'Unknown type'} is not a valid {rules['label']} Excel"
        return True, "Excel format accepted"
    if check_id == "workbook":
        try:
            ctx.book = _open_book(ctx.payload)
            ctx.sheet = ctx.book.active
        except Exception as exc:  # noqa: BLE001
            return False, str(exc) or "The workbook could not be opened"
        return True, "Workbook opened"
    if check_id == "sheet":
        if ctx.book is None:
            return False, "The workbook could not be opened"
        wanted = str((rules.get("sheet") or {}).get("name") or "").strip()
        if not wanted:
            ctx.sheet = ctx.book.active
            return True, "Sheet selected"
        match = next((name for name in ctx.book.sheetnames if name.strip().upper() == wanted.upper()), None)
        if match is None:
            return False, f"Sheet '{wanted}' was not found"
        ctx.sheet = ctx.book[match]
        return True, f"Sheet {match} found"
    if check_id == "issue_name":
        try:
            ctx.file_version = _version_from_filename(ctx.filename)
            ctx.issue_number = int(ctx.file_version.split(".", 1)[0])
        except ValueError as exc:
            return False, str(exc)
        return True, f"Issue {ctx.file_version} has been detected"
    if check_id == "issue_version":
        if not ctx.file_version:
            return False, "Filename must contain XX.XX"
        try:
            ctx.header_cells = _sheet_headers(ctx.sheet, rules)
        except ValueError as exc:
            return False, str(exc)
        wanted = ctx.file_version
        pattern = str((rules.get("issue_number") or {}).get("header_pattern") or r"^\d{2}\.\d{2}$")
        if not re.match(pattern, wanted):
            return False, f"Filename issue {wanted} does not match XX.XX"
        keys = {_header_key(name) for name in ctx.header_cells}
        aliases = {_header_key(wanted), _header_key(f"{int(wanted.split('.', 1)[0])}.{wanted.split('.', 1)[1]}")}
        if not keys & aliases:
            return False, f"Column {wanted} was not found in PD-SUMMARY"
        return True, f"Column {wanted} found"
    if check_id == "issue_number":
        spec = rules.get("issue_number") or {}
        wanted = ctx.file_version
        if not wanted:
            return False, "Filename must contain XX.XX"
        if not ctx.headers:
            try:
                ctx.headers = _map_headers(ctx.sheet, rules)
            except Exception as exc:  # noqa: BLE001
                return False, str(exc)
        if not ctx.rows:
            try:
                ctx.rows = _read_rows(ctx.sheet, rules, ctx.headers)
            except Exception as exc:  # noqa: BLE001
                return False, str(exc)
        col = next((key for key in (ctx.rows[0] if ctx.rows else {}) if _version_id(key) == wanted), None)
        if col is None:
            return False, f"Column {wanted} was not found in PD-SUMMARY"
        allowed = [str(item) for item in spec.get("enum") or ["Y", "OUT", "UNDER REVIEW"]]
        allowed_norm = {item.upper() for item in allowed}
        for row in ctx.rows:
            value = (row.get(col) or "").strip()
            if not spec.get("allow_empty", False) and not value:
                detail = f"Row {row['_line']}: {wanted} is empty"
                _attach_fail_rows(ctx, detail)
                return False, detail
            if value and value.upper() not in allowed_norm:
                detail = f"Row {row['_line']}: {wanted} must be {' / '.join(allowed)}"
                _attach_fail_rows(ctx, detail)
                return False, detail
        return True, f"Column {wanted} values are valid"
    if check_id == "headers":
        try:
            ctx.headers = _map_headers(ctx.sheet, rules)
        except Exception as exc:  # noqa: BLE001
            return False, str(exc)
        return True, "Required columns found"
    if check_id == "empty":
        try:
            ctx.rows = _read_rows(ctx.sheet, rules, ctx.headers)
        except Exception as exc:  # noqa: BLE001
            return False, str(exc)
        errors = _empty_errors(rules, ctx.rows)
        return _fail_detail(ctx, errors, "No empty required cells")
    if check_id == "values":
        return _fail_detail(ctx, _value_errors(rules, ctx.rows), "Column formats are valid")
    if check_id == "flags":
        return _fail_detail(ctx, _flag_errors(rules, ctx.rows), "Flags are sequential")
    if check_id == "unique":
        return _fail_detail(ctx, _unique_errors(rules, ctx.rows), "References are unique")
    if check_id == "unique_line":
        errors, fail_rows = _line_unique_errors(ctx.rows, ctx.file_version)
        seen = {item.get("line") for item in ctx.fail_rows}
        for snap in fail_rows:
            if snap.get("line") not in seen:
                ctx.fail_rows.append(snap)
                seen.add(snap.get("line"))
        return (False, _joined(errors)) if errors else (True, f"No duplicated lines on issue {ctx.file_version or 'XX.XX'}")
    if check_id == "store":
        if ctx.failed:
            return False, "Not stored because other checks failed"
        return (True, "Ready to store") if ctx.rows else (False, "No rows to store")
    return False, f"Unknown check {check_id}"


def to_pd_rows(rows: list[dict], rules: dict) -> tuple[list[str], list[dict]]:
    versions = _version_ids(rules, rows)
    ids = [column["id"] for column in rules["columns"]]
    parsed = []
    for row in rows:
        item = {col: row.get(col) or "" for col in ids}
        item["flags"] = {col: row.get(col) or "" for col in versions}
        parsed.append(item)
    return versions, parsed


def to_named_rows(rows: list[dict], rules: dict) -> list[dict]:
    ids = [column["id"] for column in rules["columns"]]
    return [{col: row.get(col) or "" for col in ids} for row in rows]


def write_sample_workbook(path: Path, rows: list[dict]) -> None:
    rules = load_rules("pd")
    book = Workbook()
    sheet = book.active
    sheet.title = str((rules.get("sheet") or {}).get("name") or "PD-SUMMARY")
    versions = list(rows[0]["flags"].keys()) if rows else []
    columns = list(rules["columns"])
    sheet.append([column["header"] for column in columns] + versions)
    for cell in sheet[1]:
        cell.number_format = "@"
    for row in rows:
        line = [row.get(column["id"], "") for column in columns]
        line.extend(row["flags"][col] for col in versions)
        sheet.append(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
